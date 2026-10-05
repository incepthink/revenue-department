"""Drive the Phase 4 screens end to end in a browser:
  Drafts     officer opens an analysed docket's drafts, regenerates all four, reviews the internal render (provenance
             markers) and the issue render (none), downloads the issue .docx, edits it, uploads it (diff shown), rates it
  Clauses    officer curates: mined candidates grouped by key / cluster, live preview, a template with a literal
             number is refused
  Backtest   admin opens the latest run: gate banner, per-holdout table, adjudicates a false positive and labels a
             precedent's relevance — the label-based metrics update

Needs: `make api`, `make web`, `make users`, the synthetic pool with holdouts + a backtest run, an analysed case, and
Playwright (`pip install playwright`; PW_CHROMIUM=<chrome.exe> reuses an installed browser).

    python eval/e2e/e2e_phase4.py DKT-002

Screenshots go to var/e2e_shots/phase4_*.png.
"""
import os
import re
import sys
from pathlib import Path

from docx import Document
from playwright.sync_api import expect, sync_playwright

BASE = "http://localhost:3000"
SHOTS = Path(__file__).resolve().parents[2] / "var" / "e2e_shots"
SHOTS.mkdir(parents=True, exist_ok=True)
PW = "dev-password-1"
CASE_REF = sys.argv[1] if len(sys.argv) > 1 else "DKT-002"
log = lambda *a: print(*a, flush=True)  # noqa: E731


def login(ctx, email):
    p = ctx.new_page()
    p.set_viewport_size({"width": 1500, "height": 950})
    p.goto(f"{BASE}/login")
    p.fill("#email", email)
    p.fill("#password", PW)
    p.click("button:has-text('Sign in')")
    expect(p.locator("h1", has_text="Cases")).to_be_visible()
    return p


def drafts(p):
    row = p.locator("tr", has=p.locator("td >> nth=0", has_text=re.compile(rf"^{re.escape(CASE_REF)}$")))
    row.locator("a", has_text=re.compile(CASE_REF)).first.click()
    p.locator("a.btn", has_text="Drafts").click()
    expect(p.locator("h1", has_text="Drafts")).to_be_visible()
    p.locator("button.primary", has_text=re.compile("all four")).click()
    expect(p.locator("button.primary", has_text=re.compile("Regenerate all four"))).to_be_enabled(timeout=60000)
    cards = p.locator("section.card", has=p.locator("h2"))
    expect(cards).to_have_count(4)
    for name in ("छाननी टिप्पणी", "त्रुटी पत्र", "मागणी पत्र", "जिल्हाधिकारी आदेश"):
        card = cards.filter(has=p.locator("h2", has_text=name))
        expect(card.locator(".badge", has_text="leak check passed")).to_be_visible()
        expect(card.locator("a.btn", has_text="Issue .docx")).to_be_visible()
    log("drafts: four templates generated, every issue render passed the leak check")
    q = cards.filter(has=p.locator("h2", has_text="त्रुटी पत्र"))
    q.locator("button", has_text="Review").click()
    body = q.locator(".grid2 .card").first
    expect(body).to_contain_text("[टीप", timeout=20000)
    expect(body).to_contain_text("आढळून आल्या आहेत")
    p.screenshot(path=SHOTS / "phase4_01_drafts_internal.png", full_page=True)
    q.locator(".tabs button", has_text="Issue").click()
    expect(body).to_contain_text("त्रुटी पत्र")
    assert "[टीप" not in body.inner_text()
    with p.expect_download() as dl:
        q.locator("a.btn", has_text="Issue .docx").click()
    doc = Document(dl.value.path())
    first = next(par for par in doc.paragraphs if par.text.startswith("१. "))
    for r in first.runs[1:]:
        r.text = ""
    first.runs[0].text = "१. मिळकत पत्रिकेची प्रमाणित प्रत सादर करावी."
    doc.add_paragraph("६. अधिकाऱ्यांनी जोडलेला मुद्दा.")
    edited = SHOTS / "edited_query_letter.docx"
    doc.save(edited)
    q.locator("input[type=file]").set_input_files(str(edited))
    expect(q.locator(".notice", has_text="Edited version stored")).to_be_visible(timeout=20000)
    expect(q.locator(".badge", has_text=re.compile(r"similarity \d+%"))).to_be_visible()
    q.locator("button[aria-label='4 of 5']").click()
    q.locator("button", has_text="Save rating").click()
    expect(q.locator(".notice", has_text="Rating saved")).to_be_visible()
    p.screenshot(path=SHOTS / "phase4_02_drafts_issue_diff.png", full_page=True)
    log("drafts: issue downloaded, edited, uploaded (diff shown), rated")


def clauses(p):
    p.goto(f"{BASE}/admin/clauses")
    expect(p.locator("h1", has_text="Clause library")).to_be_visible()
    p.locator(".tabs button", has_text="त्रुटी पत्र").click()
    item = p.locator("section.card", has=p.locator(".mono", has_text=re.compile(r"^query\.closing$")))
    cand = item.locator("div[style*='cursor']", has=p.locator(".badge", has_text="mined")).first
    cand.click()
    editor = p.locator("section.card.stack")
    expect(editor.locator("textarea")).to_be_visible()
    expect(editor.locator("b", has_text="Preview")).to_be_visible()
    expect(editor).to_contain_text("Found in")
    editor.locator("textarea").fill("तरी वरील त्रुटींची पूर्तता ३० दिवसांत करावी.")
    expect(editor.locator(".error", has_text="literal digits")).to_be_visible(timeout=10000)
    editor.locator("button", has_text="Make canonical").click()
    expect(editor.locator("ul.error", has_text="literal digits").last).to_be_visible()
    editor.locator("textarea").fill("तरी वरील त्रुटींची पूर्तता {{ reply_days|mr_words }} दिवसांत करावी.")
    expect(editor.locator("div", has_text=re.compile("तीस दिवसांत")).last).to_be_visible(timeout=10000)
    p.screenshot(path=SHOTS / "phase4_03_clauses.png", full_page=True)
    log("clauses: mined candidate opened, literal number refused, slot preview rendered (तीस दिवसांत)")


def backtest(ctx):
    p = login(ctx, "admin@rci.local")
    p.goto(f"{BASE}/backtest")
    expect(p.locator("h1", has_text="Backtest")).to_be_visible()
    runs = p.locator("section", has=p.locator("h2", has_text="Runs")).locator("tbody tr")
    expect(runs.first).to_be_visible()
    holdouts = p.locator("section", has=p.locator("h2", has_text="Holdouts")).locator("tbody tr")
    expect(holdouts).to_have_count(20)
    runs.first.locator("a").click()
    expect(p.locator("h1", has_text="Backtest — synthetic")).to_be_visible()
    expect(p.locator("b", has_text=re.compile(r"Gate (PASS|FAIL)"))).to_be_visible()
    p.screenshot(path=SHOTS / "phase4_04_backtest_summary.png", full_page=True)
    p.locator(".tabs button", has_text="Per holdout").click()
    expect(p.locator("table.data tbody tr")).to_have_count(20)
    p.locator(".tabs button", has_text="False-positive").click()
    first = p.locator("table.data tbody tr").first
    expect(first).to_be_visible()
    first.locator("button", has_text="valid but not asked").click()
    expect(first.locator("button.ok", has_text="valid but not asked")).to_be_visible()
    p.screenshot(path=SHOTS / "phase4_05_adjudication.png", full_page=True)
    p.locator(".tabs button", has_text="Precedent relevance").click()
    sec = p.locator("section.card", has=p.locator("tbody tr")).first  # the earliest holdouts have no earlier precedent
    sec.locator("tbody tr").first.locator("button", has_text="2").click()
    expect(sec.locator("tbody tr").first.locator("button.primary", has_text="2")).to_be_visible()
    p.locator(".tabs button", has_text="Summary").click()
    prec = p.locator("section.card", has=p.locator("h3", has_text="Precedents"))
    expect(prec.locator("tr", has_text="P@5 (officer)").locator("b")).not_to_have_text("—")
    expect(p.locator("section.card", has=p.locator("h3", has_text="Queries")).locator("tr", has_text="adjudicated precision").locator("b")).not_to_have_text("—")
    rep = p.request.get(f"{BASE}/api/rci/backtest/runs/{p.url.rsplit('/', 1)[-1]}/report.html")
    assert rep.ok and "Per holdout" in rep.text()
    p.screenshot(path=SHOTS / "phase4_06_backtest_labels.png", full_page=True)
    log("backtest: run opened, gate shown, FP adjudicated + relevance labelled → metrics updated, HTML report served")


def main():
    with sync_playwright() as pw:
        browser = pw.chromium.launch(executable_path=os.environ.get("PW_CHROMIUM") or None)
        ctx = browser.new_context(accept_downloads=True)
        p = login(ctx, "officer@rci.local")
        drafts(p)
        clauses(p)
        backtest(browser.new_context())
        browser.close()
    log("E2E phase 4: OK")


if __name__ == "__main__":
    main()
