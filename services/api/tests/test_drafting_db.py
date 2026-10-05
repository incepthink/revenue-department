"""Drafting end-to-end on the test DB: seeded clause library, mining + curation through the API, all four drafts for an
analysed incoming docket and for a decided precedent (both render modes, stored blobs, versions, roles), the issue
download detokenising the case's own PII through the audited vault door, officer upload + redacted diff, ratings."""
import io
import uuid

import pytest
from docx import Document
from sqlalchemy import func, select

from conftest import ingest_case
from rci.db.models import AuditLog, Case, CaseCard, Clause, Draft, EvalLabel
from rci.drafting import render
from rci.drafting import service as drafting
from rci.intel import service as intel
from rci.intel import synthetic
from rci.pipeline.redact.vault import SqlVaultStore, Vault
from rci.seed import _load
from rci.storage import get_blob_store

pytestmark = pytest.mark.db
PREFIX = "DR-P"
N = 12
TEMPLATES = {"SCRUTINY_NOTE", "QUERY_LETTER", "DEMAND_LETTER", "COLLECTOR_ORDER"}


@pytest.fixture(scope="module")
def world(db_url):
    from rci.db.session import session_factory, session_scope

    with session_scope(db_url) as s:
        synthetic.generate_pool(s, N, seed=5, prefix=PREFIX)
    s = session_factory(db_url)()
    incoming = ingest_case(s, "DRF-DKT-002", "DKT-002")
    intel.analyse_case(s, incoming, explain_llm=False)
    granted = next(c for c in s.scalars(select(Case).where(Case.case_ref.like(f"{PREFIX}-%")).order_by(Case.case_ref))
                   if s.scalar(select(CaseCard).where(CaseCard.case_id == c.id)).card["facts"]["payments"])
    intel.analyse_case(s, granted, explain_llm=False)
    s.commit()
    out = {"incoming_id": incoming.id, "granted_id": granted.id}
    s.close()
    return out


def test_seed_clauses_are_loaded_once(db_session):
    from rci.drafting.clauses import seed_clauses

    n = len(_load("clauses.yaml")["clauses"])
    count = lambda: db_session.scalar(select(func.count()).select_from(Clause).where(Clause.origin == "seed"))  # noqa: E731
    assert count() == n
    seed_clauses(db_session)
    assert count() == n
    assert db_session.scalar(select(func.count()).select_from(Clause).where(Clause.origin == "seed", Clause.status == "canonical",
                                                                           Clause.verification == "to_verify")) == n


def test_generate_view_download(api, world, db_session):
    officer, reviewer, anon = api("officer"), api("reviewer1"), api(None)
    cid = world["incoming_id"]
    assert anon.post(f"/cases/{cid}/drafts").status_code == 401
    assert reviewer.post(f"/cases/{cid}/drafts").status_code == 403
    r = officer.post(f"/cases/{cid}/drafts")
    assert r.status_code == 201, r.text
    got = {d["template_code"]: d for d in r.json()}
    assert set(got) == TEMPLATES
    for code, d in got.items():
        assert d["issue"]["leak_check"]["ok"] and d["issue"]["downloadable"] and d["internal"]["downloadable"], code
        assert d["internal"]["version"] == d["issue"]["version"] == 1
    assert got["QUERY_LETTER"]["status"] == "generated"
    assert got["DEMAND_LETTER"]["status"] == "incomplete"  # no verified rate / ASR in the test rules layer
    assert any("premium not computed" in w for w in got["DEMAND_LETTER"]["issue"]["warnings"])
    issue = officer.get(f"/drafts/{got['QUERY_LETTER']['issue']['id']}").json()
    internal = officer.get(f"/drafts/{got['QUERY_LETTER']['internal']['id']}").json()
    assert issue["paragraphs"] and not issue["notes"] and internal["notes"]
    assert any(p.startswith("१. ") for p in issue["paragraphs"])
    precedent_refs = {n for n in " ".join(x["text"] for x in internal["notes"]).split() if n.startswith("TST-P") or n.startswith(PREFIX)}
    assert precedent_refs and not any(ref.rstrip(",") in " ".join(issue["paragraphs"]) for ref in precedent_refs)
    assert reviewer.get(f"/drafts/{issue['id']}").status_code == 403
    for mode in ("issue", "internal"):
        dl = officer.get(f"/drafts/{got['COLLECTOR_ORDER'][mode]['id']}/download")
        assert dl.status_code == 200 and dl.headers["content-type"].startswith(render.DOCX_MIME)
        text = render.docx_all_text(dl.content)
        assert ("AI-सहाय्यित मसुदा" in text) == (mode == "internal")
    listing = officer.get(f"/cases/{cid}/drafts").json()
    assert {x["template_code"] for x in listing["latest"]} == TEMPLATES and len(listing["history"]) == 8
    again = {d["template_code"]: d for d in officer.post(f"/cases/{cid}/drafts", json={"templates": ["QUERY_LETTER"]}).json()}
    assert again["QUERY_LETTER"]["version"] == 2
    db_session.expire_all()
    old = db_session.get(Draft, uuid.UUID(got["QUERY_LETTER"]["issue"]["id"]))
    assert old.status == "superseded"
    actions = set(db_session.scalars(select(AuditLog.action).where(AuditLog.case_id == cid)))
    assert {"draft.generate", "draft.download", "auth.denied"} <= actions


def test_decided_case_order_is_complete(db_session, world):
    case = db_session.get(Case, world["granted_id"])
    res = {r["template_code"]: r for r in drafting.generate(db_session, case)}
    db_session.commit()
    order = res["COLLECTOR_ORDER"]["issue"]
    assert res["COLLECTOR_ORDER"]["status"] == "generated" and not order["missing_slots"], order["warnings"]
    card = db_session.scalar(select(CaseCard).where(CaseCard.case_id == case.id)).card
    grn = card["facts"]["payments"][0]["grn"]["value"]
    row = db_session.get(Draft, uuid.UUID(order["id"]))
    assert any(grn in p for p in row.content["paragraphs"])


def test_issue_download_detokenises_own_pii_only_for_officers(db_session, world, users):
    case_id = world["incoming_id"]
    vault = Vault.from_settings(SqlVaultStore(db_session))
    token = vault.tokenize(case_id, "PERSON", "श्री. राम & सन्स <अध्यक्ष>")
    inputs, _ = drafting.load_inputs(db_session, db_session.get(Case, case_id), llm=False)
    from rci.drafting.context import build

    model = build("QUERY_LETTER", inputs, "issue")
    model.context["to"] = [f"{token},"] + model.context["to"]
    data = render.render_docx("QUERY_LETTER", model.context, internal=False, cfg=inputs.cfg)
    key = f"drafts/{case_id}/TEST/pii.docx"
    get_blob_store().put("out", key, data)
    d = Draft(case_id=case_id, template_code="QUERY_LETTER", version=99, render_mode="issue", blob_key=key, status="generated",
              content={"paragraphs": render.docx_paragraphs(data)})
    db_session.add(d)
    db_session.flush()
    out, name = drafting.download(db_session, d, users["officer"])
    text = render.docx_all_text(out)
    assert "श्री. राम & सन्स <अध्यक्ष>" in text and token not in text and name.endswith("_issue.docx")
    Document(io.BytesIO(out))  # still a valid document (values XML-escaped)
    audit = db_session.scalars(select(AuditLog).where(AuditLog.action == "vault.detokenize", AuditLog.case_id == case_id)).all()
    assert audit and audit[-1].details["draft_ref"] == str(d.id) and "राम" not in str(audit[-1].details)
    with pytest.raises(PermissionError):
        drafting.download(db_session, d, users["reviewer1"])
    db_session.rollback()


def test_officer_upload_diff_and_rating(api, world, db_session):
    officer = api("officer")
    cid = world["incoming_id"]
    latest = {x["template_code"]: x for x in officer.get(f"/cases/{cid}/drafts").json()["latest"]}["QUERY_LETTER"]
    issue_id, internal_id = latest["issue"]["id"], latest["internal"]["id"]
    doc = Document(io.BytesIO(officer.get(f"/drafts/{issue_id}/download").content))
    first = next(p for p in doc.paragraphs if p.text.startswith("१. "))
    for r in first.runs[1:]:
        r.text = ""
    first.runs[0].text = "१. मिळकत पत्रिकेची प्रमाणित प्रत सादर करावी. संपर्क 9876543210"
    doc.add_paragraph("६. नवीन मुद्दा अधिकाऱ्यांनी जोडला.")
    buf = io.BytesIO()
    doc.save(buf)
    files = {"file": ("edited.docx", buf.getvalue(), render.DOCX_MIME)}
    r = officer.post(f"/drafts/{issue_id}/upload", files=files)
    assert r.status_code == 200, r.text
    diff = r.json()
    assert diff["paragraphs"]["changed"] + diff["paragraphs"]["inserted"] >= 2 and diff["clauses_edited"] >= 1
    stored = officer.get(f"/drafts/{issue_id}").json()
    assert stored["status"] == "officer_edited" and stored["officer_diff"]["ops"]
    assert "9876543210" not in str(stored["officer_diff"]) and "[PHONE_" in str(stored["officer_diff"])  # redacted before storage
    assert officer.post(f"/drafts/{internal_id}/upload", files=files).status_code == 422
    assert officer.post(f"/drafts/{issue_id}/upload", files={"file": ("x.docx", b"not a zip", "application/octet-stream")}).status_code == 422
    for rating in (4, 3):
        assert officer.post(f"/drafts/{issue_id}/rating", json={"rating": rating, "usability": "minor_edits"}).status_code == 200
    labels = db_session.scalars(select(EvalLabel).where(EvalLabel.case_id == cid, EvalLabel.label_kind == "draft_rating")).all()
    assert len(labels) == 1 and labels[0].value["rating"] == 3 and labels[0].label_key == "QUERY_LETTER"
    assert officer.post(f"/drafts/{issue_id}/rating", json={"rating": 9, "usability": "as_is"}).status_code == 422


def test_clause_mining_and_curation(api, world, db_session):
    admin, officer, reviewer = api("admin"), api("officer"), api("reviewer1")
    assert officer.post("/clauses/mine", json={"data_class": "synthetic"}).status_code == 403
    mined = admin.post("/clauses/mine", json={"data_class": "synthetic"}).json()
    assert mined["created"] > 0 and mined["by_data_class"]["synthetic"]["query_letter"]["rows"] > 0
    cands = reviewer.get("/clauses", params={"kind": "query_letter", "status": "candidate"}).json()
    items = [c for c in cands if c["clause_key"].startswith("query_item.") and c["clause_key"] != "query_item.unmapped"]
    assert items and all("{{" not in c["text_mr"] for c in items)
    closing = next(c for c in cands if c["clause_key"] == "query.closing" and not c["problems"])
    assert closing["n_sources"] >= 1 and all(r["case_ref"].startswith(("DR-P", "TST-P", "BT-P")) for r in closing["source_refs"])
    assert reviewer.patch(f"/clauses/{closing['id']}", json={"status": "canonical"}).status_code == 403
    bad = officer.patch(f"/clauses/{closing['id']}", json={"status": "canonical", "template": "तरी ३० दिवसांत पूर्तता करावी."})
    assert bad.status_code == 422 and "literal digits" in str(bad.json())
    assert officer.patch(f"/clauses/{closing['id']}", json={"verification": "verified"}).status_code == 403
    ok = officer.patch(f"/clauses/{closing['id']}", json={"status": "canonical", "conditions": {}})
    assert ok.status_code == 200 and ok.json()["origin"] == "curated" and ok.json()["status"] == "canonical"
    assert admin.patch(f"/clauses/{closing['id']}", json={"verification": "verified"}).json()["verification"] == "verified"
    prev = officer.post("/clauses/preview", json={"template": "{{ society_name }} — {{ area_sqm|mr_num }} चौ.मी.",
                                                  "clause_key": "query.intro", "case_id": str(world["incoming_id"])}).json()
    assert prev["slots_from"] == "case" and "गणेश कृपा" in prev["text"] and not prev["problems"]
    # the next draft uses the curated wording (class-specific, verified) and cites its source cases internally only
    res = {d["template_code"]: d for d in officer.post(f"/cases/{world['incoming_id']}/drafts", json={"templates": ["QUERY_LETTER"]}).json()}
    internal = officer.get(f"/drafts/{res['QUERY_LETTER']['internal']['id']}").json()
    used = {c["clause_key"]: c for c in internal["clauses_used"]}
    assert used["query.closing"]["clause_id"] == closing["id"] and used["query.closing"]["origin"] == "curated"
    assert res["QUERY_LETTER"]["issue"]["leak_check"]["ok"]
    created = officer.post("/clauses", json={"clause_key": "order.condition.extra", "kind": "collector_order", "data_class": "synthetic",
                                             "template": "सदर आदेश {{ order_date|mr_date }} पासून अंमलात येईल.", "status": "canonical"})
    assert created.status_code == 201 and created.json()["origin"] == "curated"
