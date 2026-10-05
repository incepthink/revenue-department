"""Build the four docxtpl templates (config/templates/*.docx) with python-docx.

    rci build-templates        # regenerate after changing the layout here or the font in config/drafting.yaml

The generated files are ordinary Word documents with Jinja tags; the office can restyle them in Word (keep the tags).
Every style carries the configured Unicode Devanagari font in all four rFonts slots — Devanagari is a *complex
script* in Word, so the cs slot (and bCs / szCs) is what actually applies — and the Marathi language tag.
Paragraph loops use docxtpl's {%p %} tags and table rows {%tr %} (each tag alone in its own paragraph / row).
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FILES = {"SCRUTINY_NOTE": "scrutiny_note.docx", "QUERY_LETTER": "query_letter.docx", "DEMAND_LETTER": "demand_letter.docx",
         "COLLECTOR_ORDER": "collector_order.docx"}


def set_run_font(rpr, font: str, size_pt: float | None = None, bold: bool | None = None) -> None:
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font)
    rfonts.set(qn("w:hint"), "cs")
    if bold is not None:
        for tag in ("w:b", "w:bCs"):
            el = rpr.find(qn(tag))
            if el is None:
                el = OxmlElement(tag)
                rpr.append(el)
            el.set(qn("w:val"), "1" if bold else "0")
    if size_pt:
        for tag in ("w:sz", "w:szCs"):
            el = rpr.find(qn(tag))
            if el is None:
                el = OxmlElement(tag)
                rpr.append(el)
            el.set(qn("w:val"), str(int(size_pt * 2)))
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "mr-IN")
    lang.set(qn("w:bidi"), "mr-IN")


class Doc:
    def __init__(self, font: str, size_pt: float):
        self.font, self.size = font, size_pt
        self.d = Document()
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        for style in self.d.styles:  # every built-in style (headings, table grid, …) gets the Devanagari font
            if style.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER, WD_STYLE_TYPE.TABLE):
                set_run_font(style.element.get_or_add_rPr(), font)
        normal = self.d.styles["Normal"]
        set_run_font(normal.element.get_or_add_rPr(), font, size_pt)
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.15
        rpr_default = self.d.styles.element.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}")
        if rpr_default is not None:
            set_run_font(rpr_default, font, size_pt)

    def p(self, text: str, *, bold: bool = False, align: str | None = None, size: float | None = None, color: str | None = None,
          indent_cm: float | None = None, space_after: float | None = None):
        para = self.d.add_paragraph()
        run = para.add_run(text)
        set_run_font(run._r.get_or_add_rPr(), self.font, size or self.size, bold)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if align:
            para.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                              "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
        if indent_cm is not None:
            para.paragraph_format.left_indent = Cm(indent_cm)
        if space_after is not None:
            para.paragraph_format.space_after = Pt(space_after)
        return para

    def tag(self, tag: str):
        return self.p(tag)

    def loop(self, head: str, body: str, *, bold: bool = False, align: str | None = None, size: float | None = None,
             indent_cm: float | None = None):
        self.tag("{%p " + head + " %}")
        self.p(body, bold=bold, align=align, size=size, indent_cm=indent_cm)
        self.tag("{%p endfor %}")

    def table(self, loop: str, cells: list[str], headers: list[str] | None = None, widths_cm: list[float] | None = None,
              size: float | None = None):
        rows = (1 if headers else 0) + 3
        t = self.d.add_table(rows=rows, cols=len(cells))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        r = 0
        if headers:
            for c, h in zip(t.rows[0].cells, headers):
                self._cell(c, h, bold=True, size=size)
            r = 1
        self._cell(t.rows[r].cells[0], "{%tr " + loop + " %}")
        for c, text in zip(t.rows[r + 1].cells, cells):
            self._cell(c, text, size=size)
        self._cell(t.rows[r + 2].cells[0], "{%tr endfor %}")
        if widths_cm:
            t.autofit = False
            grid = t._tbl.tblGrid
            for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
                gc.set(qn("w:w"), str(int(Cm(w).twips)))
            for row in t.rows:
                for c, w in zip(row.cells, widths_cm):
                    c.width = Cm(w)
        self.d.add_paragraph()
        return t

    def _cell(self, cell, text: str, *, bold: bool = False, size: float | None = None):
        para = cell.paragraphs[0]
        run = para.add_run(text)
        set_run_font(run._r.get_or_add_rPr(), self.font, size or self.size - 1, bold)

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        cp = self.d.core_properties
        cp.author, cp.title, cp.comments, cp.last_modified_by = "", "", "", ""
        self.d.save(str(path))


# ------------------------------------------------------------------------------------------------
# Shared blocks
# ------------------------------------------------------------------------------------------------
def letter_header(d: Doc, *, addressed: bool = True) -> None:
    d.loop("for l in letterhead", "{{ l }}", bold=True, align="center")
    d.p("क्रमांक : {{ ref_no }}", space_after=0)
    d.p("दिनांक : {{ date }}")
    d.p("{{ title }}", bold=True, align="center", size=d.size + 2)
    if addressed:
        d.p("प्रति,", space_after=0)
        d.loop("for l in to", "{{ l }}", indent_cm=0.8)
        d.p("विषय : {{ subject }}", bold=True)
        d.p("संदर्भ : {{ reference }}")


def signature(d: Doc) -> None:
    d.p("")
    d.loop("for l in signatory", "{{ l }}", align="right")


def notes(d: Doc) -> None:
    d.tag("{%p if internal and notes %}")
    d.p("{{ notes_title }}", bold=True, size=d.size - 2, color="9C2A00")
    d.loop("for n in notes", "[टीप {{ n.n_mr }}] {{ n.text }}", size=d.size - 3)
    d.tag("{%p endif %}")


# ------------------------------------------------------------------------------------------------
# The four templates
# ------------------------------------------------------------------------------------------------
def query_letter(d: Doc) -> None:
    letter_header(d)
    d.p("{{ intro }}", align="justify")
    d.loop("for it in items", "{{ it.n }}. {{ it.text }}", align="justify", indent_cm=0.6)
    d.p("{{ closing }}", align="justify")
    signature(d)
    notes(d)


def demand_letter(d: Doc) -> None:
    letter_header(d)
    d.p("{{ intro }}", align="justify")
    d.table("for l in lines", ["{{ l.label }}", "{{ l.value }}"], widths_cm=[9.5, 7.0])
    d.p("एकूण देय रक्कम : {{ total }}", bold=True)
    d.p("{{ total_words }}")
    d.p("{{ payment_instruction }}", align="justify")
    d.tag("{%p if lease_note %}")
    d.p("{{ lease_note }}", align="justify")
    d.tag("{%p endif %}")
    d.p("{{ closing }}", align="justify")
    signature(d)
    notes(d)


def collector_order(d: Doc) -> None:
    letter_header(d, addressed=False)
    d.p("वाचले :", bold=True)
    d.loop("for r in read", "{{ r.n }}) {{ r.text }}", indent_cm=0.6)
    d.loop("for w in whereas", "{{ w }}", align="justify")
    d.p("{{ operative }}", align="justify")
    d.p("अटी व शर्ती :", bold=True)
    d.loop("for c in conditions", "{{ c.n }}. {{ c.text }}", align="justify", indent_cm=0.6)
    signature(d)
    d.p("प्रत :", bold=True, space_after=0)
    d.loop("for c in copy_to", "{{ c }}", indent_cm=0.6)
    notes(d)


def scrutiny_note(d: Doc) -> None:
    d.loop("for l in letterhead", "{{ l }}", bold=True, align="center")
    d.p("{{ title }}", bold=True, align="center", size=d.size + 2)
    d.p("विषय : {{ subject }}", bold=True)
    d.p("{{ intro }}", align="justify")
    d.p("१. प्रकरणाची थोडक्यात हकीकत", bold=True)
    d.p("{{ narrative }}", align="justify")
    d.p("२. अर्जदार व जमिनीचा तपशील", bold=True)
    d.table("for r in facts", ["{{ r.label }}", "{{ r.value }}"], widths_cm=[6.0, 10.5])
    d.p("३. सादर / प्राप्त कागदपत्रे", bold=True)
    d.table("for r in documents", ["{{ r.n }}", "{{ r.name }}", "{{ r.stage }}", "{{ r.date }}", "{{ r.pages }}"],
            headers=["अ. क्र.", "कागदपत्र", "टप्पा", "दिनांक", "पृष्ठे"], widths_cm=[1.4, 7.0, 3.0, 2.8, 2.3])
    d.p("४. कागदपत्रांची तपासणी", bold=True)
    d.table("for r in checks", ["{{ r.name }}", "{{ r.result }}"], widths_cm=[10.0, 6.5])
    d.p("५. आढळलेल्या त्रुटी", bold=True)
    d.tag("{%p if no_gaps %}")
    d.p("छाननीत कोणतीही त्रुटी आढळलेली नाही.")
    d.tag("{%p endif %}")
    d.tag("{%p if not no_gaps %}")
    d.table("for r in gaps", ["{{ r.n }}", "{{ r.name }}", "{{ r.status }}", "{{ r.basis }}", "{{ r.why }}"],
            headers=["अ. क्र.", "बाब", "स्थिती", "आधार", "कारण"], widths_cm=[1.4, 3.8, 2.2, 3.4, 5.7], size=d.size - 2)
    d.tag("{%p endif %}")
    d.p("६. नजराणा परिगणना", bold=True)
    d.table("for r in fee_rows", ["{{ r.label }}", "{{ r.value }}"], widths_cm=[8.5, 8.0])
    d.p("{{ fee_status }}")
    d.tag("{%p if steps %}")
    d.p("७. पुढील कार्यवाही (समान निर्णीत प्रकरणांनुसार)", bold=True)
    d.table("for r in steps", ["{{ r.n }}", "{{ r.name }}", "{{ r.when }}"],
            headers=["अ. क्र.", "कार्यवाही", "अपेक्षित दिनांक"], widths_cm=[1.4, 10.0, 5.1])
    d.tag("{%p endif %}")
    d.p("८. प्रस्ताव", bold=True)
    d.p("{{ proposal }}", align="justify")
    d.p("{{ closing }}")
    signature(d)
    notes(d)


BUILDERS = {"SCRUTINY_NOTE": scrutiny_note, "QUERY_LETTER": query_letter, "DEMAND_LETTER": demand_letter,
            "COLLECTOR_ORDER": collector_order}


def build_all(out_dir: Path, font: str, size_pt: float = 12) -> list[Path]:
    out = []
    for code, fn in BUILDERS.items():
        d = Doc(font, size_pt)
        fn(d)
        path = out_dir / FILES[code]
        d.save(path)
        out.append(path)
    return out
