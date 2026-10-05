"""docx rendering, the internal watermark, text extraction, the issue leak check and the officer-edit diff.

    render_docx(template_code, context, internal=…, cfg=…)  → bytes (docxtpl; values XML-escaped)
    docx_paragraphs(data)                                  → body paragraphs in order (table rows as "a | b | c")
    leak_check(data, forbidden, own_token)                 → {"ok", "found"}: an issue render may not contain another
                                                             case's ref / id / exhibit id / society, the watermark,
                                                             provenance markers or a PII token that is not this case's
    diff(generated, uploaded, clauses_used)                → paragraph-level diff + which clauses the officer kept
"""
from __future__ import annotations

import difflib
import io
import zipfile
from pathlib import Path

from jinja2 import Environment
from lxml import etree

from rci.drafting.config import DraftingConfig
from rci.drafting.context import MARK_OPEN
from rci.drafting.docxbuild import FILES, build_all
from rci.drafting.mrtext import FILTERS
from rci.pipeline.normalise import ascii_digits
from rci.pipeline.redact.vault import TOKEN_RE

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W, "v": "urn:schemas-microsoft-com:vml"}
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def template_path(code: str, cfg: DraftingConfig) -> Path:
    path = cfg.templates_dir / FILES[code]
    if not path.exists():  # first run / deleted: build from code with the configured font
        build_all(cfg.templates_dir, str(cfg.value("font", "Nirmala UI")), float(cfg.value("font_size_pt", 12)))
    return path


def _jinja_env() -> Environment:
    e = Environment(autoescape=True)
    e.filters.update(FILTERS)
    return e


WATERMARK = """<w:r xmlns:w="{w}" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">
<w:rPr><w:noProof/></w:rPr><w:pict>
<v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" path="m@7,l@8,m@5,21600l@6,21600e">
<v:formulas><v:f eqn="sum #0 0 10800"/><v:f eqn="prod #0 2 1"/><v:f eqn="sum 21600 0 @1"/><v:f eqn="sum 0 0 @2"/>
<v:f eqn="sum 21600 0 @3"/><v:f eqn="if @0 @3 0"/><v:f eqn="if @0 21600 @1"/><v:f eqn="if @0 0 @2"/><v:f eqn="if @0 @4 21600"/>
<v:f eqn="mid @5 @6"/><v:f eqn="mid @8 @5"/><v:f eqn="mid @7 @8"/><v:f eqn="mid @6 @7"/><v:f eqn="sum @6 0 @5"/></v:formulas>
<v:path textpathok="t" o:connecttype="custom" o:connectlocs="@9,0;@10,10800;@11,21600;@12,10800" o:connectangles="270,180,90,0"/>
<v:textpath on="t" fitshape="t"/><o:lock v:ext="edit" text="t" shapetype="t"/></v:shapetype>
<v:shape id="RciWaterMark{n}" o:spid="_x0000_s20{n}5" type="#_x0000_t136"
 style="position:absolute;margin-left:0;margin-top:0;width:470pt;height:94pt;rotation:315;z-index:-251655168;mso-position-horizontal:center;mso-position-horizontal-relative:margin;mso-position-vertical:center;mso-position-vertical-relative:margin"
 o:allowincell="f" fillcolor="#c0c0c0" stroked="f"><v:fill opacity=".45"/>
<v:textpath style="font-family:&quot;{font}&quot;;font-size:1pt" string="{text}"/></v:shape></w:pict></w:r>"""


def add_watermark(document, text: str, header_line: str, font: str) -> None:
    """Diagonal VML text watermark (what Word's own Watermark command writes) + a red header line, on every section."""
    from docx.oxml import parse_xml
    from docx.shared import Pt, RGBColor

    from rci.drafting.docxbuild import set_run_font

    for n, section in enumerate(document.sections):
        header = section.header
        header.is_linked_to_previous = False
        para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        run = para.add_run(header_line)
        set_run_font(run._r.get_or_add_rPr(), font, 9, True)
        run.font.color.rgb = RGBColor(0xB4, 0x23, 0x18)
        run.font.size = Pt(9)
        esc = lambda s: s.replace("&", "&amp;").replace('"', "&quot;").replace("<", "&lt;")  # noqa: E731
        para._p.append(parse_xml(WATERMARK.format(w=W, n=n + 1, font=esc(font), text=esc(text))))


def render_docx(template_code: str, context: dict, *, internal: bool, cfg: DraftingConfig) -> bytes:
    from docxtpl import DocxTemplate

    tpl = DocxTemplate(str(template_path(template_code, cfg)))
    tpl.render(context, jinja_env=_jinja_env(), autoescape=True)
    doc = tpl.docx
    cp = doc.core_properties
    cp.author = cp.last_modified_by = str(context.get("letterhead", [""])[0] if context.get("letterhead") else "")
    cp.title, cp.comments, cp.keywords = str(context.get("title", "")), "", ""
    if internal:
        add_watermark(doc, str(cfg.value("watermark", "AI-सहाय्यित मसुदा")), str(cfg.value("internal_header", "")),
                      str(cfg.value("font", "Nirmala UI")))
        cp.comments = str(cfg.value("watermark", ""))
    bio = io.BytesIO()
    tpl.save(bio)
    return bio.getvalue()


# ------------------------------------------------------------------------------------------------
# Text extraction
# ------------------------------------------------------------------------------------------------
def _para_text(p) -> str:
    return "".join(t.text or "" for t in p.iter(f"{{{W}}}t"))


def docx_paragraphs(data: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    body = root.find(f"{{{W}}}body")
    out = []
    for el in body:
        if el.tag == f"{{{W}}}p":
            t = _para_text(el).strip()
            if t:
                out.append(t)
        elif el.tag == f"{{{W}}}tbl":
            for tr in el.iter(f"{{{W}}}tr"):
                cells = [" ".join(_para_text(p).strip() for p in tc.iter(f"{{{W}}}p")).strip() for tc in tr.iter(f"{{{W}}}tc")]
                if any(cells):
                    out.append(" | ".join(cells))
    return out


def docx_all_text(data: bytes) -> str:
    """Every visible or embedded string: body, headers, footers, VML watermark text and document properties."""
    chunks = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for name in z.namelist():
            if not name.endswith(".xml"):
                continue
            xml = z.read(name)
            if name.startswith("word/"):
                root = etree.fromstring(xml)
                chunks += [_para_text(p) for p in root.iter(f"{{{W}}}p")]
                chunks += [tp.get("string") or "" for tp in root.iter("{urn:schemas-microsoft-com:vml}textpath")]
            elif name.startswith("docProps/"):
                chunks.append(" ".join(etree.fromstring(xml).itertext()))
    return "\n".join(chunks)


def leak_check(data: bytes, forbidden: list[str], own_token, *, cfg: DraftingConfig) -> dict:
    text = docx_all_text(data)
    found = []
    for s in forbidden:
        if s and s in text:
            found.append({"kind": "other_case", "value": s})
    for s in (cfg.value("watermark"), cfg.value("internal_header"), MARK_OPEN):
        if s and s in text:
            found.append({"kind": "internal_marking", "value": s})
    for m in TOKEN_RE.finditer(text):
        if not own_token(m.group(0)):
            found.append({"kind": "foreign_pii_token", "value": m.group(0)})
    return {"ok": not found, "found": found}


def detokenize_docx(data: bytes, replace) -> bytes:
    """Rewrite every XML part through `replace(xml_text) -> xml_text` (vault detokenisation for the issue download)."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            blob = src.read(info)
            if info.filename.endswith(".xml") and TOKEN_RE.search(blob.decode("utf-8", "ignore")):
                blob = replace(blob.decode("utf-8")).encode("utf-8")
            z.writestr(info, blob)
    return out.getvalue()


# ------------------------------------------------------------------------------------------------
# Officer edits
# ------------------------------------------------------------------------------------------------
def _numbers(paras: list[str]) -> set[str]:
    """Dates (dd/mm/yyyy) as one token, then the remaining numbers without grouping commas."""
    from rci.drafting.crosscheck import DATE_DMY, NUMBER

    out = set()
    for p in paras:
        t = ascii_digits(p)
        out |= {m.group(0) for m in DATE_DMY.finditer(t)}
        out |= {n.replace(",", "") for n in NUMBER.findall(DATE_DMY.sub(" ", t))}
    return out


def diff(generated: list[str], uploaded: list[str], clauses_used: list[dict] | None = None) -> dict:
    sm = difflib.SequenceMatcher(a=generated, b=uploaded, autojunk=False)
    ops, counts = [], {"equal": 0, "replace": 0, "insert": 0, "delete": 0}
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        counts[tag] += max(i2 - i1, j2 - j1)
        if tag != "equal":
            a, b = generated[i1:i2], uploaded[j1:j2]
            ops.append({"op": tag, "generated": a, "uploaded": b,
                        "similarity": round(difflib.SequenceMatcher(a="\n".join(a), b="\n".join(b)).ratio(), 3)})
    joined_a, joined_b = "\n".join(generated), "\n".join(uploaded)
    up_text = joined_b
    clauses = [{"clause_id": c["clause_id"], "clause_key": c["clause_key"], "kept": c["text"] in up_text}
               for c in clauses_used or [] if c.get("text") and "________" not in c["text"]]
    na, nb = _numbers(generated), _numbers(uploaded)
    return {
        "paragraphs": {"generated": len(generated), "uploaded": len(uploaded), "unchanged": counts["equal"],
                       "changed": counts["replace"], "inserted": counts["insert"], "deleted": counts["delete"]},
        "similarity": round(difflib.SequenceMatcher(a=joined_a, b=joined_b, autojunk=False).ratio(), 4),
        "numbers_removed": sorted(na - nb), "numbers_added": sorted(nb - na),
        "clauses": clauses, "clauses_edited": sum(1 for c in clauses if not c["kept"]),
        "ops": ops[:200],
    }
